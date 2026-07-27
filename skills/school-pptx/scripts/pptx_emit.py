#!/usr/bin/env python3
"""Mechanical emitter for immutable school-pptx physical deck plans."""

from __future__ import annotations

from importlib import import_module
from io import BytesIO
import hashlib
import os
import posixpath
from pathlib import Path, PurePosixPath
from typing import Any, BinaryIO, Callable
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile, ZipInfo

from pptx_model import PHYSICAL_PLAN_SCHEMA, MediaBinding, PhysicalDeckPlan, PhysicalSlide
from pptx_objects import (
    PptxObjectError,
    MediaReference,
    add_bound_contain_picture,
    add_bound_gallery_card,
    add_contain_picture,
    add_fragment_text_frame,
    add_gallery_card,
    add_literal_text,
    add_rich_text,
    add_table,
    add_timeline,
    set_notes,
)
from pptx_ooxml import remove_seed_slides


class PptxEmitError(RuntimeError):
    """Bounded renderer-domain failure safe for CLI presentation."""

    def __init__(self, code: str, message: str):
        super().__init__(f"{code}: {message}")
        self.code = code
        self.message = message


def normalize_package_metadata(target: Path | BinaryIO) -> None:
    """Rewrite ZIP container metadata so identical OOXML has identical bytes."""
    if hasattr(target, "read"):
        target.seek(0)
        payload = target.read()
    else:
        payload = Path(target).read_bytes()
    source_buffer = BytesIO(payload)
    normalized_buffer = BytesIO()
    with ZipFile(source_buffer, "r") as source, ZipFile(normalized_buffer, "w") as normalized:
        normalized.comment = source.comment
        for original in source.infolist():
            info = ZipInfo(original.filename, (1980, 1, 1, 0, 0, 0))
            info.compress_type = original.compress_type
            info.comment = original.comment
            info.external_attr = original.external_attr
            info.internal_attr = original.internal_attr
            info.create_system = original.create_system
            normalized.writestr(info, source.read(original.filename))
    normalized_payload = normalized_buffer.getvalue()
    if hasattr(target, "write"):
        target.seek(0)
        target.truncate(0)
        target.write(normalized_payload)
        target.flush()
        os.fsync(target.fileno())
        target.seek(0)
    else:
        Path(target).write_bytes(normalized_payload)


MAX_PACKAGE_ENTRIES = 512
MAX_PACKAGE_ENTRY_BYTES = 16 * 1024 * 1024
MAX_PACKAGE_TOTAL_BYTES = 96 * 1024 * 1024
MAX_PACKAGE_RELATIONSHIPS = 4096
FORBIDDEN_XML_MARKERS = (b"<!DOCTYPE", b"<!ENTITY")
REL_NS = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
OBJECT_ERROR_MESSAGES = {
    "PPTX_MEDIA_SIZE_LIMIT": "媒体文件大小超过安全上限。",
    "PPTX_MEDIA_FORMAT_INVALID": "媒体格式不受支持，请使用 JPEG 或 PNG。",
    "PPTX_MEDIA_PIXEL_LIMIT": "媒体像素数量超过安全上限。",
    "PPTX_MEDIA_PATH_INVALID": "媒体路径包含不安全的目录项。",
    "PPTX_MEDIA_HASH_INVALID": "冻结媒体哈希不匹配。",
    "PPTX_TABLE_ROW_HEIGHT_MISMATCH": "冻结表格行高与表格行数不一致。",
}
VISUAL_LAYOUTS = frozenset({
    # Phase 43's historic image-text/gallery plans remain accepted without a
    # VisualPlan.  These four layouts were introduced with the Phase 51 bound
    # planning contract and must always carry it.
    "hero-evidence", "annotated-image", "process-overlay", "kpi-evidence",
})


def _plan_error(message: str) -> PptxEmitError:
    return PptxEmitError("PPTX_FROZEN_PLAN_INVALID", message)


def _is_sha256(value: object) -> bool:
    return isinstance(value, str) and len(value) == 64 and all(character in "0123456789abcdef" for character in value)


def _validate_media_binding(binding: MediaBinding) -> None:
    if (not binding.confirmed or not _is_sha256(binding.sha256) or binding.payload is None
            or len(binding.payload) != binding.byte_count):
        raise _plan_error("冻结媒体绑定不完整。")
    if hashlib.sha256(binding.payload).hexdigest() != binding.sha256:
        raise _plan_error("冻结媒体哈希不匹配。")
    expected_format = "png" if binding.payload.startswith(b"\x89PNG\r\n\x1a\n") else "jpeg" if binding.payload.startswith(b"\xff\xd8\xff") else None
    if expected_format != binding.media_format:
        raise _plan_error("冻结媒体格式不匹配。")


def _validate_frozen_plan(plan: PhysicalDeckPlan, manifest: dict[str, Any]) -> bool:
    """Validate a complete Phase 51 plan without deriving any visual decision.

    Returns whether this is a bound Phase 51 plan.  Unbound plans containing only
    historic layouts remain supported for the Phase 43 compatibility contract.
    """
    if not isinstance(plan, PhysicalDeckPlan):
        raise _plan_error("发射器只接受 PhysicalDeckPlan。")
    bound_values = (plan.visual_plan_sha256, plan.markdown_sha256, plan.manifest_sha256, plan.template_sha256)
    bound = any(value is not None for value in bound_values)
    if not bound:
        if any(slide.layout in VISUAL_LAYOUTS for slide in plan.slides):
            raise _plan_error("视觉布局必须使用完整冻结计划。")
        return False
    if plan.schema_version != PHYSICAL_PLAN_SCHEMA or not all(_is_sha256(value) for value in bound_values):
        raise _plan_error("冻结计划 schema 或输入哈希无效。")
    if not _is_sha256(plan.integrity_sha256) or plan.integrity_sha256 != plan.expected_integrity_sha256:
        raise _plan_error("冻结物理计划完整性绑定无效。")
    expected_indices = tuple(range(len(plan.slides)))
    if tuple(slide.physical_index for slide in plan.slides) != expected_indices:
        raise _plan_error("冻结物理页索引不连续。")
    seen: list[int] = []
    for logical_index, physical_indices in plan.logical_to_physical:
        if logical_index < 0 or not physical_indices:
            raise _plan_error("冻结 logical-to-physical 映射不完整。")
        seen.extend(physical_indices)
        if any(index < 0 or index >= len(plan.slides) or plan.slides[index].logical_index != logical_index
               for index in physical_indices):
            raise _plan_error("冻结 logical-to-physical 映射不匹配。")
    mapped = tuple(sorted(seen))
    non_closing = tuple(index for index, slide in enumerate(plan.slides) if slide.layout != "closing")
    if mapped != non_closing or len(seen) != len(set(seen)):
        raise _plan_error("冻结 logical-to-physical 映射未完整覆盖。")
    for slide in plan.slides:
        layout_manifest = manifest.get("layouts", {}).get(slide.layout)
        if not isinstance(layout_manifest, dict):
            raise _plan_error("冻结计划包含未知布局。")
        if slide.layout not in VISUAL_LAYOUTS:
            continue
        expected_roles = tuple(str(role) for role in layout_manifest.get("object_roles", ()))
        if slide.object_roles != expected_roles:
            raise _plan_error("冻结对象角色与 manifest 不匹配。")
        if not slide.variant or slide.variant != slide.preset_id:
            raise _plan_error("冻结视觉变体不完整。")
        try:
            variant_layout, count_text = slide.variant.rsplit(":", 1)
            count = int(count_text)
        except (AttributeError, ValueError):
            raise _plan_error("冻结视觉变体无效。") from None
        if variant_layout != slide.layout or count != len(slide.atomic_groups) or count != len(slide.fragments):
            raise _plan_error("冻结视觉变体与原子组不匹配。")
        if any(not group or not group[0].startswith("id=") for group in slide.atomic_groups):
            raise _plan_error("冻结原子组不完整。")
        expected_media = {
            "hero-evidence": (1, 1),
            "annotated-image": (1, 1),
            "process-overlay": (0, 1),
            "kpi-evidence": (0, 0),
        }.get(slide.layout)
        if expected_media is not None and not expected_media[0] <= len(slide.media_bindings) <= expected_media[1]:
            raise _plan_error("冻结视觉页媒体绑定数量无效。")
        for binding in slide.media_bindings:
            if binding.placement != "contain":
                raise _plan_error("冻结视觉页媒体 placement 无效。")
            _validate_media_binding(binding)
    return True


def map_object_error(error: PptxObjectError) -> PptxEmitError:
    code = str(error) if str(error) in OBJECT_ERROR_MESSAGES else "PPTX_OBJECT_INVALID"
    return PptxEmitError(code, OBJECT_ERROR_MESSAGES.get(code, "PPTX 对象内容无效。"))


def _frozen_body_typography(physical: Any, slot_id: str) -> dict[str, float]:
    fonts = dict(physical.selected_font_sizes)
    if slot_id not in fonts:
        raise PptxEmitError("PPTX_PLAN_FONT_MISSING", "冻结物理计划缺少正文槽字号。")
    values = dict(physical.slot_values)
    keys = (
        "margin_left_points", "margin_right_points", "margin_top_points", "margin_bottom_points",
        "line_spacing", "paragraph_spacing_points",
    )
    try:
        typography = {key: float(values[f"typography.{slot_id}.{key}"]) for key in keys}
    except (KeyError, TypeError, ValueError) as exc:
        raise PptxEmitError("PPTX_PLAN_TYPOGRAPHY_MISSING", "冻结物理计划缺少正文槽排版参数。") from exc
    typography["font_size"] = float(fonts[slot_id])
    return typography


def _frozen_code_typography(physical: Any) -> dict[str, float]:
    return _frozen_body_typography(physical, "code")


def _frozen_font(physical: Any, key: str) -> float:
    try:
        value = dict(physical.selected_font_sizes)[key]
        return float(value)
    except (KeyError, TypeError, ValueError) as exc:
        raise PptxEmitError("PPTX_PLAN_FONT_MISSING", "冻结物理计划缺少发射所需字号。") from exc


def _slot(layout: dict[str, Any], slot_id: str) -> dict[str, Any]:
    slot = next((item for item in layout.get("slots", ()) if item.get("id") == slot_id), None)
    if not isinstance(slot, dict) or not isinstance(slot.get("geometry"), dict):
        raise _plan_error(f"视觉布局缺少 manifest 槽位 {slot_id}。")
    return slot


def _visual_binding(physical: PhysicalSlide, index: int = 0) -> MediaBinding:
    try:
        binding = physical.media_bindings[index]
    except IndexError as exc:
        raise _plan_error("冻结视觉页缺少媒体绑定。") from exc
    _validate_media_binding(binding)
    return binding


def _bound_picture(slide: Any, binding: MediaBinding, geometry: dict[str, int], name: str) -> Any:
    # _validate_media_binding has already established all fields; retain explicit
    # checks here so this helper remains fail-closed when called independently.
    if binding.payload is None or binding.sha256 is None or binding.media_format is None:
        raise _plan_error("冻结媒体绑定不完整。")
    return add_bound_contain_picture(
        slide, binding.payload, geometry, expected_sha256=binding.sha256,
        expected_format=binding.media_format, name=name,
    )


def _add_card(slide: Any, geometry: dict[str, int], *, name: str) -> Any:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, geometry["x"], geometry["y"], geometry["width"], geometry["height"])
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(242, 246, 250)
    shape.line.color.rgb = RGBColor(172, 190, 207)
    return shape


def _inset(geometry: dict[str, int], top: float, height: float) -> dict[str, int]:
    return {
        "x": geometry["x"] + int(geometry["width"] * 0.08),
        "y": geometry["y"] + int(geometry["height"] * top),
        "width": int(geometry["width"] * 0.84),
        "height": max(1, int(geometry["height"] * height)),
    }


def _card_text_region(
    geometry: dict[str, int], *, left: float, top: float, width: float, height: float,
) -> dict[str, int]:
    """Return a manifest-relative KPI text region with a fixed card inset."""
    return {
        "x": geometry["x"] + int(geometry["width"] * left),
        "y": geometry["y"] + int(geometry["height"] * top),
        "width": max(1, int(geometry["width"] * width)),
        "height": max(1, int(geometry["height"] * height)),
    }


def _kpi_values(fragment: Any) -> tuple[str, str, str, str]:
    row = dict(fragment.metadata).get("row", "")
    values = tuple(part.strip() for part in row.split("|"))
    if len(values) != 4 or not all(values):
        raise _plan_error("冻结 KPI 原子组不完整。")
    return values[0], values[1], values[2], values[3]


def _emit_kpi_cards(slide: Any, physical: PhysicalSlide, layout: dict[str, Any], highlight: str) -> set[str]:
    cards_slot = _slot(layout, "kpi_cards")
    count = len(physical.fragments)
    presets = cards_slot.get("item_presets", {}).get(count)
    if not isinstance(presets, list) or len(presets) != count:
        raise _plan_error("冻结 KPI 变体未对应 manifest 预设。")
    roles: set[str] = set()
    minimum = _frozen_font(physical, "font.kpi_cards.min")
    maximum = _frozen_font(physical, "font.kpi_cards.max")
    for index, (fragment, preset) in enumerate(zip(physical.fragments, presets), start=1):
        geometry = preset.get("card") if isinstance(preset, dict) else None
        if not isinstance(geometry, dict):
            raise _plan_error("KPI 卡片几何无效。")
        label, value, unit, footnote = _kpi_values(fragment)
        _add_card(slide, geometry, name=f"school-pptx:kpi-card:{index}")
        add_rich_text(slide, _card_text_region(geometry, left=.08, top=.08, width=.84, height=.20), label,
                      font_size=minimum, highlight_scheme=highlight,
                      name=f"school-pptx:kpi-label:{index}")
        add_literal_text(slide, _card_text_region(geometry, left=.08, top=.32, width=.38, height=.40), value,
                         font_size=maximum,
                         name=f"school-pptx:kpi-value:{index}", font_name="Aptos Display")
        add_literal_text(slide, _card_text_region(geometry, left=.48, top=.50, width=.22, height=.24), unit,
                         font_size=minimum,
                         name=f"school-pptx:kpi-unit:{index}")
        add_rich_text(slide, _card_text_region(geometry, left=.08, top=.76, width=.84, height=.18), footnote,
                      font_size=minimum, highlight_scheme=highlight,
                      name=f"school-pptx:kpi-footnote:{index}")
        roles.update({"kpi-card", "kpi-value", "kpi-unit", "kpi-footnote", "kpi-label"})
    return roles


def _emit_hero_evidence(slide: Any, physical: PhysicalSlide, layout: dict[str, Any], highlight: str) -> set[str]:
    binding = _visual_binding(physical)
    media = _slot(layout, "hero_media")
    _bound_picture(slide, binding, media["geometry"], "school-pptx:hero-picture")
    add_rich_text(slide, _inset(media["geometry"], .88, .10), binding.caption, font_size=_frozen_font(physical, "font.hero_media.min"),
                  highlight_scheme=highlight, name="school-pptx:hero-caption")
    return {"hero-picture", "caption"} | _emit_kpi_cards(slide, physical, layout, highlight)


def _emit_annotated_image(slide: Any, physical: PhysicalSlide, layout: dict[str, Any], highlight: str) -> set[str]:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE

    binding = _visual_binding(physical)
    media = _slot(layout, "annotated_media")
    annotations = _slot(layout, "annotations")
    anchors = annotations.get("fixed_anchors", {})
    _bound_picture(slide, binding, media["geometry"], "school-pptx:annotated-picture")
    add_rich_text(slide, _inset(media["geometry"], .88, .10), binding.caption, font_size=_frozen_font(physical, "font.annotated_media.min"),
                  highlight_scheme=highlight, name="school-pptx:annotated-caption")
    for index, fragment in enumerate(physical.fragments, start=1):
        anchor = anchors.get(index) or anchors.get(str(index))
        label = dict(fragment.metadata).get("label")
        if not isinstance(anchor, dict) or not isinstance(label, str) or not label:
            raise _plan_error("冻结注解锚点或标签不完整。")
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, anchor["x"], anchor["y"], anchor["height"], anchor["height"])
        badge.name = f"school-pptx:annotation-number:{index}"
        badge.fill.solid(); badge.fill.fore_color.rgb = RGBColor(0, 112, 192)
        badge.text_frame.paragraphs[0].text = str(index)
        label_geometry = {"x": anchor["x"] + anchor["width"], "y": anchor["y"],
                          "width": max(1, annotations["geometry"]["x"] + annotations["geometry"]["width"] - anchor["x"] - anchor["width"]),
                          "height": anchor["height"]}
        add_rich_text(slide, label_geometry, label, font_size=_frozen_font(physical, "font.annotations.min"), highlight_scheme=highlight,
                      name=f"school-pptx:annotation-label:{index}")
        leader = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                            anchor["x"] + anchor["height"] // 2, anchor["y"] + anchor["height"] // 2,
                                            anchor["x"] + anchor["width"], anchor["y"] + anchor["height"] // 2)
        leader.name = f"school-pptx:annotation-leader:{index}"
    return {"annotated-picture", "caption", "annotation-number", "annotation-label", "annotation-leader", "fixed-anchor"}


def _emit_process_overlay(slide: Any, physical: PhysicalSlide, layout: dict[str, Any], highlight: str) -> set[str]:
    from pptx.enum.shapes import MSO_CONNECTOR

    process = _slot(layout, "process_nodes")
    count = len(physical.fragments)
    if count < 2 or count > 6:
        raise _plan_error("冻结流程节点数量无效。")
    region = process["geometry"]
    gap = max(1, region["width"] // (count * 10))
    node_width = (region["width"] - gap * (count - 1)) // count
    node_height = max(1, int(region["height"] * .70))
    node_y = region["y"] + (region["height"] - node_height) // 2
    for index, fragment in enumerate(physical.fragments, start=1):
        row = tuple(part.strip() for part in dict(fragment.metadata).get("row", "").split("|"))
        if len(row) < 2 or not all(row[:2]):
            raise _plan_error("冻结流程节点内容不完整。")
        geometry = {"x": region["x"] + (index - 1) * (node_width + gap), "y": node_y,
                    "width": node_width, "height": node_height}
        _add_card(slide, geometry, name=f"school-pptx:process-node:{index}")
        add_literal_text(slide, _inset(geometry, .12, .20), str(index), font_size=_frozen_font(physical, "font.process_nodes.max"),
                         name=f"school-pptx:step-number:{index}")
        add_rich_text(slide, _inset(geometry, .40, .42), "\n".join(row[:2]), font_size=_frozen_font(physical, "font.process_nodes.min"),
                      highlight_scheme=highlight, name=f"school-pptx:step-label:{index}")
        if index > 1:
            connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                geometry["x"] - gap, node_y + node_height // 2, geometry["x"], node_y + node_height // 2)
            connector.name = f"school-pptx:process-connector:{index - 1}"
    roles = {"process-node", "process-connector", "step-number", "step-label"}
    if physical.media_bindings:
        binding = _visual_binding(physical)
        media = _slot(layout, "evidence_media")
        _bound_picture(slide, binding, media["geometry"], "school-pptx:optional-evidence-picture")
        add_rich_text(slide, _inset(media["geometry"], .80, .18), binding.caption, font_size=_frozen_font(physical, "font.evidence_media.min"),
                      highlight_scheme=highlight, name="school-pptx:process-caption")
        roles.update({"optional-evidence-picture", "caption"})
    return roles


def _emit_visual_layout(slide: Any, physical: PhysicalSlide, layout: dict[str, Any], highlight: str) -> None:
    if physical.layout == "hero-evidence":
        emitted = _emit_hero_evidence(slide, physical, layout, highlight)
    elif physical.layout == "annotated-image":
        emitted = _emit_annotated_image(slide, physical, layout, highlight)
    elif physical.layout == "process-overlay":
        emitted = _emit_process_overlay(slide, physical, layout, highlight)
    elif physical.layout == "kpi-evidence":
        emitted = _emit_kpi_cards(slide, physical, layout, highlight)
    else:
        raise _plan_error("未知冻结视觉布局。")
    required = set(physical.object_roles)
    # The only intentionally absent roles are the manifest-declared optional
    # evidence image/caption on a process page with no bound media.
    optional = {"optional-evidence-picture", "caption"} if physical.layout == "process-overlay" and not physical.media_bindings else set()
    if not required - optional <= emitted:
        raise _plan_error("冻结对象角色没有原生对象映射。")


def require_dependencies(importer: Callable[[str], Any] = import_module) -> dict[str, Any]:
    try:
        pptx = importer("pptx")
        pillow = importer("PIL")
    except (ImportError, ModuleNotFoundError) as exc:
        raise PptxEmitError("PPTX_DEPENDENCY_MISSING", "需要 python-pptx 1.0.2+ 与 Pillow。") from exc
    try:
        version = tuple(int(item) for item in str(pptx.__version__).split(".")[:2])
    except (AttributeError, ValueError) as exc:
        raise PptxEmitError("PPTX_DEPENDENCY_MISSING", "无法确认 python-pptx 版本。") from exc
    if version < (1, 0) or version >= (2, 0):
        raise PptxEmitError("PPTX_DEPENDENCY_MISSING", "python-pptx 版本必须为 >=1.0.2,<2。")
    patch = int(str(pptx.__version__).split(".")[2].split("+")[0])
    if version == (1, 0) and patch < 2:
        raise PptxEmitError("PPTX_DEPENDENCY_MISSING", "python-pptx 版本必须为 >=1.0.2,<2。")
    return {"pptx": pptx, "PIL": pillow}


def layout_part_map(presentation: Any) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for layout in presentation.slide_layouts:
        part_path = str(layout.part.partname).lstrip("/")
        if part_path in result:
            raise PptxEmitError("PPTX_TEMPLATE_LAYOUT_INVALID", "模板包含重复 layout part path。")
        result[part_path] = layout
    return result


def resolve_layouts(presentation: Any, manifest: dict[str, Any]) -> dict[str, Any]:
    by_part = layout_part_map(presentation)
    resolved: dict[str, Any] = {}
    for layout_id, specification in manifest.get("layouts", {}).items():
        part_path = specification.get("pptx_layout")
        if not isinstance(part_path, str) or part_path not in by_part:
            raise PptxEmitError(
                "PPTX_TEMPLATE_LAYOUT_INVALID", f'布局 "{layout_id}" 的 PPTX part path 不存在。'
            )
        resolved[str(layout_id)] = by_part[part_path]
    return resolved


def bootstrap_template(template_path: Path, manifest: dict[str, Any]) -> tuple[Any, dict[str, Any], dict[str, int]]:
    dependencies = require_dependencies()
    try:
        presentation = dependencies["pptx"].Presentation(str(template_path))
    except (OSError, ValueError) as exc:
        raise PptxEmitError("PPTX_TEMPLATE_INVALID", "无法打开受控 PPTX 模板。") from exc
    layouts = resolve_layouts(presentation, manifest)
    seed_count = len(presentation.slides)
    removed, notes_relationships = remove_seed_slides(presentation)
    template_contract = manifest.get("template", {})
    expected_seed_count = template_contract.get("seed_slides")
    expected_notes_relationships = template_contract.get("seed_notes_relationships")
    if (
        not isinstance(expected_seed_count, int)
        or not isinstance(expected_notes_relationships, int)
        or seed_count != expected_seed_count
        or removed != expected_seed_count
        or notes_relationships != expected_notes_relationships
    ):
        raise PptxEmitError(
            "PPTX_TEMPLATE_SEED_INVALID",
            "受控模板的 seed 或 notes relationship 数量与 manifest 不一致。",
        )
    return presentation, layouts, {
        "seed_slides": seed_count,
        "removed_seed_slides": removed,
        "removed_seed_notes_relationships": notes_relationships,
    }


def emit_deck(
    plan: PhysicalDeckPlan, manifest: dict[str, Any], template_path: Path, output: Path | BinaryIO,
    *, media_root: Path | None = None, manifest_bytes: bytes | None = None, markdown_sha256: str | None = None,
) -> dict[str, Any]:
    """Emit only the frozen plan; pagination is intentionally absent from this module."""
    bound = _validate_frozen_plan(plan, manifest)
    if bound:
        # Identity belongs to planning.  The emitter verifies it once and never
        # substitutes a path, reopens authored media, or selects a fallback variant.
        if manifest_bytes is None or markdown_sha256 is None:
            raise _plan_error("冻结计划缺少调用方输入绑定。")
        if (hashlib.sha256(manifest_bytes).hexdigest() != plan.manifest_sha256
                or markdown_sha256 != plan.markdown_sha256):
            raise _plan_error("冻结计划输入哈希不匹配。")
        try:
            template_sha256 = hashlib.sha256(Path(template_path).read_bytes()).hexdigest()
        except OSError as exc:
            raise PptxEmitError("PPTX_TEMPLATE_INVALID", "无法读取受控 PPTX 模板。") from exc
        if template_sha256 != plan.template_sha256:
            raise _plan_error("冻结计划模板哈希不匹配。")
    presentation, layouts, evidence = bootstrap_template(template_path, manifest)
    highlight = manifest["inline_styles"]["highlight"]["scheme_color"]
    runtime_diagnostics: list[dict[str, Any]] = []

    def emit_font(physical: PhysicalSlide, key: str, legacy_default: float) -> float:
        """Read a frozen font, with a manifest default for an unbound legacy plan."""
        if bound or key in dict(physical.selected_font_sizes):
            return _frozen_font(physical, key)
        return legacy_default

    def record_missing(physical: Any) -> None:
        runtime_diagnostics.append({
            "code": "MEDIA_MISSING",
            "severity": "error",
            "message": "发射时媒体不存在，已保留可编辑占位符。",
            "fix": "恢复媒体文件后重新渲染。",
            "logical_indices": [physical.logical_index],
            "physical_indices": [physical.physical_index],
            "source_line": physical.source_line,
        })
    try:
        for physical in plan.slides:
            slide = presentation.slides.add_slide(layouts[physical.layout])
            layout_manifest = manifest["layouts"][physical.layout]
            slots = {slot["id"]: slot for slot in layout_manifest.get("slots", ())}
            if physical.layout == "cover":
                for slot_id, value in physical.slot_values:
                    slot = slots[slot_id]
                    font_key = f"font.{slot_id}.max"
                    font_size = (
                        _frozen_font(physical, font_key)
                        if bound or font_key in dict(physical.selected_font_sizes)
                        else float(slot["text_budget"]["font_size_max"])
                    )
                    add_rich_text(
                        slide, slot["geometry"], value,
                        font_size=font_size, highlight_scheme=highlight,
                        font_name=slot.get("font_family"),
                        font_theme_color=slot.get("font_theme_color"),
                        paragraph_alignment=slot.get("paragraph_alignment"),
                        vertical_anchor=slot.get("vertical_anchor"),
                        name=f"school-pptx:cover-{slot_id}",
                    )
                set_notes(slide, physical.notes_intent)
                continue
            if "title" in slots:
                title_font = (
                    _frozen_font(physical, "font.title.max")
                    if bound or "font.title.max" in dict(physical.selected_font_sizes)
                    else float(slots["title"]["text_budget"]["font_size_max"])
                )
                add_rich_text(
                    slide, slots["title"]["geometry"], physical.title,
                    font_size=title_font, highlight_scheme=highlight,
                    font_name=slots["title"].get("font_family"),
                    font_theme_color=slots["title"].get("font_theme_color"),
                    paragraph_alignment=slots["title"].get("paragraph_alignment"),
                    vertical_anchor=slots["title"].get("vertical_anchor"),
                    name="school-pptx:slide-title",
                )
            if physical.layout == "closing":
                set_notes(slide, physical.notes_intent)
                continue
            if physical.layout in {"hero-evidence", "annotated-image", "process-overlay", "kpi-evidence"}:
                _emit_visual_layout(slide, physical, layout_manifest, highlight)
            elif physical.layout == "code" and physical.fragments:
                code_fragments = tuple(item for item in physical.fragments if item.kind == "code")
                fragment = code_fragments[0]
                code_text = "\n".join(item.text or "" for item in code_fragments)
                code_slot = slots["code"]
                code_geometry = dict(code_slot["geometry"])
                typography = _frozen_code_typography(physical)
                if fragment.heading:
                    try:
                        heading_height = int(dict(physical.slot_values)["geometry.code.heading_height_emu"])
                    except (KeyError, TypeError, ValueError) as exc:
                        raise PptxEmitError(
                            "PPTX_PLAN_TYPOGRAPHY_MISSING",
                            "冻结物理计划缺少代码标题几何。",
                        ) from exc
                    if heading_height <= 0 or heading_height >= code_geometry["height"]:
                        raise PptxEmitError("PPTX_PLAN_TYPOGRAPHY_MISSING", "冻结代码标题几何无效。")
                    heading_geometry = {**code_geometry, "height": heading_height}
                    add_rich_text(
                        slide, heading_geometry, fragment.heading,
                        font_size=typography["font_size"], highlight_scheme=highlight,
                        margin_left_points=typography["margin_left_points"],
                        margin_right_points=typography["margin_right_points"],
                        margin_top_points=typography["margin_top_points"],
                        margin_bottom_points=typography["margin_bottom_points"],
                        line_spacing=typography["line_spacing"],
                        paragraph_spacing_points=typography["paragraph_spacing_points"],
                        name="school-pptx:code-heading",
                    )
                    code_geometry["y"] += heading_height
                    code_geometry["height"] -= heading_height
                add_literal_text(
                    slide, code_geometry, code_text,
                    font_size=typography["font_size"],
                    margin_left_points=typography["margin_left_points"],
                    margin_right_points=typography["margin_right_points"],
                    margin_top_points=typography["margin_top_points"],
                    margin_bottom_points=typography["margin_bottom_points"],
                    line_spacing=typography["line_spacing"],
                    paragraph_spacing_points=typography["paragraph_spacing_points"],
                    name="school-pptx:code",
                )
            elif physical.layout == "table" and physical.fragments:
                fragment = physical.fragments[0]
                add_table(
                    slide, slots["table"], fragment,
                    font_size=_frozen_font(physical, "table"),
                    table_name_font_size=emit_font(
                        physical, "font.table.subregions.table_name.max",
                        float(slots["table"]["subregions"]["table_name"]["text_budget"]["font_size_max"]),
                    ),
                    row_heights_emu=fragment.row_heights_emu,
                )
            elif physical.layout == "timeline" and physical.fragments:
                add_timeline(
                    slide, slots["timeline_items"], physical.fragments[0].rows, highlight_scheme=highlight,
                    font_sizes={
                        role: emit_font(
                            physical, f"font.timeline_items.node_template.subregions.{role}.min",
                            float(slots["timeline_items"]["node_template"]["subregions"][role]["text_budget"]["font_size_min"]),
                        )
                        for role in ("time", "title", "description")
                    },
                )
            elif physical.layout == "gallery":
                gallery_slot = slots["gallery_items"]
                count = len(physical.fragments)
                presets = gallery_slot["item_presets"][count]
                for index, (fragment, preset) in enumerate(zip(physical.fragments, presets)):
                    metadata = dict(fragment.metadata)
                    if bound:
                        try:
                            binding = physical.media_bindings[int(metadata["media_index"])]
                        except (KeyError, ValueError, IndexError) as exc:
                            raise _plan_error("冻结 gallery 卡片未绑定有效媒体。") from exc
                        _validate_media_binding(binding)
                        if binding.payload is None or binding.sha256 is None or binding.media_format is None:
                            raise _plan_error("冻结 gallery 媒体绑定不完整。")
                        add_bound_gallery_card(
                            slide, preset, binding.payload, metadata.get("caption", ""),
                            expected_sha256=binding.sha256, expected_format=binding.media_format, index=index,
                            highlight_scheme=highlight, font_size=emit_font(
                                physical, "font.caption.min", float(slots["caption"]["text_budget"]["font_size_min"]),
                            ),
                            paragraph_alignment=slots["caption"].get("paragraph_alignment"),
                        )
                    else:
                        _, missing, _ = add_gallery_card(
                            slide, preset, _media_path(metadata.get("authored_path", ""), media_root),
                            metadata.get("caption", ""), index=index, highlight_scheme=highlight,
                            font_size=emit_font(
                                physical, "font.caption.min", float(slots["caption"]["text_budget"]["font_size_min"]),
                            ),
                            paragraph_alignment=slots["caption"].get("paragraph_alignment"),
                        )
                        if missing:
                            record_missing(physical)
            else:
                text_fragments = [
                    fragment for fragment in physical.fragments
                    if fragment.kind != "image"
                    and not (bound and physical.layout == "image-text" and "media_index" in dict(fragment.metadata))
                ]
                target_slot_ids = tuple(dict.fromkeys(
                    fragment.target_slot or "body" for fragment in text_fragments
                ))
                if physical.layout == "two-column":
                    target_slot_ids = ("left_body", "right_body")
                for body_slot_id in target_slot_ids:
                    body_slot = slots.get(body_slot_id)
                    if body_slot is None:
                        continue
                    slot_fragments = tuple(
                        fragment for fragment in text_fragments
                        if (fragment.target_slot or "body") == body_slot_id
                    )
                    typography = _frozen_body_typography(physical, body_slot_id)
                    add_fragment_text_frame(
                        slide, body_slot["geometry"], slot_fragments,
                        font_size=typography["font_size"],
                        margin_left_points=typography["margin_left_points"],
                        margin_right_points=typography["margin_right_points"],
                        margin_top_points=typography["margin_top_points"],
                        margin_bottom_points=typography["margin_bottom_points"],
                        line_spacing=typography["line_spacing"],
                        paragraph_spacing_points=typography["paragraph_spacing_points"],
                        paragraph_alignment=body_slot.get("paragraph_alignment"),
                        vertical_anchor=body_slot.get("vertical_anchor"),
                        highlight_scheme=highlight, name=f"school-pptx:{body_slot_id}",
                    )
                image_fragment = next(
                    (
                        fragment for fragment in physical.fragments
                        if fragment.kind == "image"
                        or (bound and physical.layout == "image-text" and "media_index" in dict(fragment.metadata))
                    ),
                    None,
                )
                if image_fragment is not None and ("media" in slots or "body" in slots):
                    metadata = dict(image_fragment.metadata)
                    media_geometry = slots["media"]["geometry"] if "media" in slots else slots["body"]["geometry"]
                    if bound and "media_index" in metadata:
                        try:
                            binding = physical.media_bindings[int(metadata["media_index"])]
                        except (KeyError, ValueError, IndexError) as exc:
                            raise _plan_error("冻结 image-text 页面未绑定有效媒体。") from exc
                        _bound_picture(slide, binding, media_geometry, "school-pptx:image-text-picture")
                    else:
                        _, missing, _ = add_contain_picture(
                            slide, _media_path(metadata.get("authored_path", ""), media_root), media_geometry,
                            name="school-pptx:image-text-picture",
                        )
                        if missing:
                            record_missing(physical)
            set_notes(slide, physical.notes_intent)
    except PptxObjectError as exc:
        raise map_object_error(exc) from None
    try:
        if hasattr(output, "write"):
            output.seek(0)
            output.truncate(0)
            presentation.save(output)
            output.flush()
            os.fsync(output.fileno())
            output.seek(0)
        else:
            presentation.save(str(output))
        normalize_package_metadata(output)
    except (OSError, ValueError) as exc:
        raise PptxEmitError("PPTX_PACKAGE_INVALID", "无法保存 staged PPTX。") from exc
    package = validate_staged_package(output, expected_slides=len(plan.slides))
    return {
        **evidence,
        "physical_slides": len(plan.slides),
        "output": str(output) if isinstance(output, Path) else "<staged-stream>",
        "transition_mode": "none",
        "package": package,
        "runtime_diagnostics": runtime_diagnostics,
    }


def _media_path(authored: str, media_root: Path | None) -> MediaReference | None:
    if not authored:
        return None
    path = Path(authored)
    if path.is_absolute():
        parts = tuple(part for part in path.parts if part != path.anchor)
        if not parts or any(part in {"", ".", ".."} for part in parts):
            raise PptxObjectError("PPTX_MEDIA_PATH_INVALID")
        return MediaReference(Path(path.anchor), parts, True)
    if media_root is None:
        return None
    parts = path.parts
    if not parts or any(part in {"", ".", ".."} for part in parts):
        raise PptxObjectError("PPTX_MEDIA_PATH_INVALID")
    return MediaReference(media_root, tuple(parts), False)


def _relationship_source(name: str) -> str:
    pure = PurePosixPath(name)
    if pure.name == ".rels":
        return ""
    source_name = pure.name.removesuffix(".rels")
    parent = pure.parent.parent
    return str(parent / source_name)


def _reader(source: Path | BinaryIO) -> tuple[BinaryIO | Path, Callable[[], None]]:
    if not hasattr(source, "read"):
        return source, lambda: None
    source.flush()
    duplicate = os.fdopen(os.dup(source.fileno()), "rb")
    duplicate.seek(0)
    return duplicate, duplicate.close


def validate_staged_package(source: Path | BinaryIO, *, expected_slides: int | None = None) -> dict[str, int]:
    """Apply bounded ZIP/XML/relationship checks before a staged file may be published."""
    dependencies = require_dependencies()
    try:
        zip_source, close_zip_source = _reader(source)
        with ZipFile(zip_source) as package:
            infos = package.infolist()
            if len(infos) > MAX_PACKAGE_ENTRIES:
                raise PptxEmitError("PPTX_PACKAGE_INVALID", "PPTX ZIP entry 数量超限。")
            names = {info.filename for info in infos}
            if len(names) != len(infos):
                raise PptxEmitError("PPTX_PACKAGE_INVALID", "PPTX ZIP 包含重复 entry。")
            required = {"[Content_Types].xml", "ppt/presentation.xml", "ppt/_rels/presentation.xml.rels"}
            if not required <= names:
                raise PptxEmitError("PPTX_PACKAGE_INVALID", "PPTX 缺少核心 content type 或 presentation part。")
            total = 0
            relationship_count = 0
            slide_count = sum(
                name.startswith("ppt/slides/slide") and name.endswith(".xml") and "/_rels/" not in name
                for name in names
            )
            for info in infos:
                pure = PurePosixPath(info.filename)
                if pure.is_absolute() or ".." in pure.parts or info.file_size > MAX_PACKAGE_ENTRY_BYTES:
                    raise PptxEmitError("PPTX_PACKAGE_INVALID", "PPTX ZIP 路径或 entry 大小无效。")
                total += info.file_size
                if total > MAX_PACKAGE_TOTAL_BYTES:
                    raise PptxEmitError("PPTX_PACKAGE_INVALID", "PPTX 解压总大小超限。")
                if not (info.filename.endswith(".xml") or info.filename.endswith(".rels")):
                    continue
                payload = package.read(info)
                upper = payload.upper()
                if any(marker in upper for marker in FORBIDDEN_XML_MARKERS):
                    raise PptxEmitError("PPTX_PACKAGE_INVALID", "PPTX XML 包含外部实体声明。")
                try:
                    root = ET.fromstring(payload)
                except ET.ParseError as exc:
                    raise PptxEmitError("PPTX_PACKAGE_INVALID", "PPTX XML 无法安全解析。") from exc
                if info.filename.endswith(".rels"):
                    relationships = root.findall("r:Relationship", REL_NS)
                    relationship_count += len(relationships)
                    if relationship_count > MAX_PACKAGE_RELATIONSHIPS:
                        raise PptxEmitError("PPTX_PACKAGE_INVALID", "PPTX relationship 数量超限。")
                    relationship_source = _relationship_source(info.filename)
                    source_dir = posixpath.dirname(relationship_source)
                    for relationship in relationships:
                        if relationship.get("TargetMode") == "External":
                            raise PptxEmitError("PPTX_PACKAGE_INVALID", "PPTX 禁止 external relationship。")
                        target = relationship.get("Target", "")
                        if not target or target.startswith("/"):
                            raise PptxEmitError("PPTX_PACKAGE_INVALID", "PPTX relationship target 无效。")
                        resolved = posixpath.normpath(posixpath.join(source_dir, target))
                        if resolved == ".." or resolved.startswith("../") or resolved not in names:
                            raise PptxEmitError("PPTX_PACKAGE_INVALID", "PPTX relationship target 越界或缺失。")
            if expected_slides is not None and slide_count != expected_slides:
                raise PptxEmitError("PPTX_PACKAGE_INVALID", "PPTX slide count 与冻结计划不一致。")
    except (BadZipFile, OSError) as exc:
        raise PptxEmitError("PPTX_PACKAGE_INVALID", "PPTX ZIP 无法读取。") from exc
    finally:
        if "close_zip_source" in locals():
            close_zip_source()
    try:
        reopen_source, close_reopen_source = _reader(source)
        reopened = dependencies["pptx"].Presentation(reopen_source)
    except (OSError, ValueError, KeyError) as exc:
        raise PptxEmitError("PPTX_PACKAGE_INVALID", "PPTX 无法由 python-pptx 重开。") from exc
    finally:
        if "close_reopen_source" in locals():
            close_reopen_source()
    if expected_slides is not None and len(reopened.slides) != expected_slides:
        raise PptxEmitError("PPTX_PACKAGE_INVALID", "重开后的 slide count 与冻结计划不一致。")
    return {
        "zip_entries": len(infos),
        "uncompressed_bytes": total,
        "relationships": relationship_count,
        "slides": slide_count,
    }
